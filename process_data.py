#!/usr/bin/env python3
"""
نفس منطق app.js لكن على السيرفر (GitHub Actions).

المدخلات:
  STR Loss.xlsx            (جذر المستودع)  أو templates/Structure.xlsx
  data/UTL.xlsx            Hold + Other + AUX + ACW  (UL_lo / UL_Date)
  data/IR.xlsx             assigned_to / added_by / added_on
  data/Compensation.xlsx   Comp_ID / Comp_Da / Comp_Du
  data/Schedule.xlsx       شيت Final بعد الماكرو (ID/Date/Duration/Code) أو التقرير الخام (B=ID, C=تاريخ/كود, J=Duration)
  config.json              teleSchFactor / secondsPerTicket / scheduleCodes

المخرجات:
  output/Final_Report.xlsx  (شيتين: Matrix بنفس شكل الإكسيل + Daily_Long)

المعادلة:
  Tele-SCH  = مجموع الأكواد المختارة × 90%   (بيظهر بعد الخصم)
  Loss Time = Tele-SCH − (System + Talk Time + Comp)   (لا يقل عن صفر)
  System    = TKT × 90 ثانية  (= 0.00104166666666667 يوم)
"""
import json
import math
import re
import sys
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUTPUT = ROOT / "output" / "Final_Report.xlsx"
STRUCTURE_CANDIDATES = [ROOT / "STR Loss.xlsx", ROOT / "templates" / "Structure.xlsx"]

DEFAULT_CONFIG = {"teleSchFactor": 0.9, "secondsPerTicket": 90, "scheduleCodes": []}

ALIASES = {
    "structureId": ["Teleopti ID", "Teleopti", "ST_ID"],
    "loginId": ["Login ID", "Login", "UL_lo", "Username"],
    "perm": ["Perm"],
    "ttsUser": ["TTS User", "TTS"],
    "bssUser": ["BSS User", "BSS"],
    "group": ["Group"],
    "agentName": ["Agent Name", "Agent"],
    "status": ["Status"],
    "tlId": ["TL ID", "TL Id"],
    "tlName": ["TL Name", "Team Leader"],
    "structureDate": ["ST_D"],
    "structureDuration": ["ST_Du", "ST Duration"],
    "irAssigned": ["assigned_to"],
    "irAdded": ["added_by", "IR_L_E"],
    "irDate": ["added_on", "Date"],
    "utlUser": ["UL_lo", "Login ID", "Login"],
    "utlDate": ["UL_Date", "Date"],
    "compUser": ["User Name", "Username", "Comp_User", "TTS User"],
    "compId": ["ID", "Comp_ID", "Comp ID", "Teleopti ID", "ST_ID"],
    "compDate": ["Date", "Comp_Da"],
    "compDuration": ["Code Time", "Comp_Du", "Comp Duration"],
}
IR_COL = {"assigned": 23, "added": 24, "date": 25}  # X, Y, Z (احتياطي)
SCHEDULE_COL = {"id": 1, "date_or_code": 2, "duration": 9}  # B, C, J
TOTALS_WORDS = {"totals", "total", "الاجمالي", "الإجمالي", "اجمالي", "إجمالي", "المجموع"}
MONTHS = {m: i + 1 for i, m in enumerate("jan feb mar apr may jun jul aug sep oct nov dec".split())}
PARTIAL_MIN = 5


# ─── أدوات ────────────────────────────────────────────────────────────────
def is_blank(v):
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    try:
        return bool(pd.isna(v)) if not isinstance(v, (list, tuple, dict)) else False
    except (TypeError, ValueError):
        return False


def norm(v):
    if is_blank(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"[^a-z0-9\u0600-\u06ff]+", "", str(v).strip().lower())


def find_col(columns, aliases):
    cols = list(columns)
    wanted = [norm(a) for a in aliases if norm(a)]
    for w in wanted:
        for c in cols:
            if norm(c) == w:
                return c
    for w in wanted:
        if len(w) < PARTIAL_MIN:
            continue
        for c in cols:
            if w in norm(c):
                return c
    return None


def text(v):
    return "" if is_blank(v) else str(v).strip()


def parse_seconds(v):
    if is_blank(v) or v == "":
        return 0
    if isinstance(v, datetime):
        return v.hour * 3600 + v.minute * 60 + v.second
    if isinstance(v, time):
        return v.hour * 3600 + v.minute * 60 + v.second
    if isinstance(v, timedelta):
        return round(v.total_seconds())
    if isinstance(v, (int, float)):
        return round(v * 86400) if 0 < v < 1 else v
    s = str(v).strip()
    if not s or s.lower() in {"unpaid", "maternity", "planned sick"}:
        return 0
    try:
        parts = [float(p) for p in s.split(":")]
    except ValueError:
        return 0
    if len(parts) == 3:
        return parts[0] * 3600 + parts[1] * 60 + parts[2]
    if len(parts) == 2:
        return parts[0] * 3600 + parts[1] * 60
    return parts[0] if len(parts) == 1 else 0


def fmt_time(seconds):
    if not seconds or seconds <= 0:
        return "0:00:00"
    total = int(round(seconds))
    return f"{total // 3600}:{(total % 3600) // 60:02d}:{total % 60:02d}"


def date_key(v, default_year=None):
    """Month-First (9/1/2026 = 1 سبتمبر). أي نص مش تاريخ صريح يرجع ''."""
    if is_blank(v) or v == "":
        return ""
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        if 20000 < v < 80000:
            return (datetime(1899, 12, 30) + timedelta(days=float(v))).date().isoformat()
        return ""
    s = str(v).strip()
    if re.fullmatch(r"\d{5}(?:\.\d+)?", s):
        return date_key(float(s))
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[T\s].*)?$", s)
    if m:
        return f"{int(m[1]):04d}-{int(m[2]):02d}-{int(m[3]):02d}"
    m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})(?:\s.*)?$", s)
    if m:
        a, b, y = int(m[1]), int(m[2]), int(m[3])
        y = y + 2000 if y < 100 else y
        month, day = (b, a) if (a > 12 and b <= 12) else (a, b)
        if not (1 <= month <= 12 and 1 <= day <= 31):
            return ""
        return f"{y:04d}-{month:02d}-{day:02d}"
    m = re.match(r"^(\d{1,2})[\s\-/]([A-Za-z]{3})[A-Za-z]*(?:[\s\-/,]+(\d{2,4}))?(?:\s.*)?$", s)
    if m:
        month = MONTHS.get(m[2].lower())
        if not month:
            return ""
        y = int(m[3]) if m[3] else (default_year or datetime.now().year)
        y = y + 2000 if y < 100 else y
        return f"{y:04d}-{month:02d}-{int(m[1]):02d}"
    m = re.search(r"\b(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})\b", s)
    if m:
        return date_key(f"{m[1]}/{m[2]}/{m[3]}")
    mon = r"(Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|June?|July?|Aug(?:ust)?|Sept?(?:ember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)"
    m = re.search(rf"\b{mon}\.?\s+(\d{{1,2}})(?:st|nd|rd|th)?,?\s+(\d{{4}})\b", s, re.I)
    if m:
        return f"{int(m[3]):04d}-{MONTHS[m[1][:3].lower()]:02d}-{int(m[2]):02d}"
    m = re.search(rf"\b(\d{{1,2}})(?:st|nd|rd|th)?\s+{mon}\.?,?\s+(\d{{4}})\b", s, re.I)
    if m:
        return f"{int(m[3]):04d}-{MONTHS[m[2][:3].lower()]:02d}-{int(m[1]):02d}"
    return ""


def lookup_key(value, day):
    n = norm(value)
    return f"{n}|{day}" if n and day else ""


def add(idx, value, day, amount):
    k = lookup_key(value, day)
    if k:
        idx[k] = idx.get(k, 0) + amount


def lookup_first(idx, candidates, day):
    for c in candidates:
        k = lookup_key(c, day)
        if k and k in idx:
            return idx[k]
    return 0


# ─── القراءة ──────────────────────────────────────────────────────────────
def find_data_file(stem):
    for ext in (".xlsx", ".xls", ".csv"):
        p = DATA / f"{stem}{ext}"
        if p.exists():
            return p
    return None


def read_sheets(path, header=0):
    if path.suffix.lower() == ".csv":
        return {"csv": pd.read_csv(path, header=header, dtype=object)}
    return pd.read_excel(path, sheet_name=None, header=header, dtype=object)


def rows_by_header(path, hints, groups):
    """يدوّر على صف العناوين (أول 40 صف) اللي فيه كل المجموعات المطلوبة."""
    sheets = read_sheets(path, header=None)
    names = sorted(sheets, key=lambda n: 0 if any(h in norm(n) for h in hints) else 1)
    for name in names:
        df = sheets[name]
        for i in range(min(len(df), 40)):
            cells = [norm(c) for c in df.iloc[i].tolist()]
            if all(any(norm(a) in cells for a in g) for g in groups):
                headers, seen = [], {}
                for j, h in enumerate(df.iloc[i].tolist()):
                    base = text(h) or f"Column {j + 1}"
                    n = seen.get(base, 0)
                    seen[base] = n + 1
                    headers.append(f"{base} {n + 1}" if n else base)
                body = df.iloc[i + 1:].dropna(how="all")
                body.columns = headers
                return body.to_dict("records")
    return None


def read_rows(stem, hints, groups=None, validate=None):
    path = find_data_file(stem)
    if not path:
        return []
    if groups:
        found = rows_by_header(path, hints, groups)
        if found:
            return found
    sheets = read_sheets(path)
    names = sorted(sheets, key=lambda n: 0 if any(h in norm(n) for h in hints) else 1)
    for name in names:
        df = sheets[name].dropna(how="all")
        if df.empty:
            continue
        rows = df.to_dict("records")
        if not validate or validate(rows):
            return rows
    return []


def read_structure():
    for path in STRUCTURE_CANDIDATES:
        if path.exists():
            sheets = read_sheets(path)
            names = sorted(sheets, key=lambda n: 0 if any(h in norm(n) for h in ("sep", "updated", "structure")) else 1)
            for name in names:
                rows = sheets[name].dropna(how="all").to_dict("records")
                if rows and find_col(rows[0].keys(), ALIASES["agentName"]) is not None:
                    return rows
    return []


def val(row, key, default=""):
    col = find_col(row.keys(), ALIASES[key])
    return row[col] if col is not None else default


# ─── الفهارس ──────────────────────────────────────────────────────────────
def sum_index(rows, user_key, date_key_name, value_key):
    idx = {}
    for r in rows:
        add(idx, val(r, user_key), date_key(val(r, date_key_name)), parse_seconds(val(r, value_key, 0)))
    return idx


def talk_index(rows):
    idx = {}
    for r in rows:
        def g(*names):
            c = find_col(r.keys(), list(names))
            return parse_seconds(r[c]) if c is not None else 0

        total = g("Hold Time", "HoldTime") + g("Other Time", "OtherTime") + g("AUXOUTOFFTIME") + g("ACWOUTOFFTIME")
        add(idx, val(r, "utlUser"), date_key(val(r, "utlDate")), total)
    return idx


def user_id(v):
    return text(v).split("@")[0]


def ir_indexes(rows):
    assigning, tkt = {}, {}
    if not rows:
        return assigning, tkt
    keys = list(rows[0].keys())
    c_assigned = find_col(keys, ALIASES["irAssigned"]) or (keys[IR_COL["assigned"]] if len(keys) > IR_COL["assigned"] else None)
    c_added = find_col(keys, ALIASES["irAdded"]) or (keys[IR_COL["added"]] if len(keys) > IR_COL["added"] else None)
    c_date = find_col(keys, ALIASES["irDate"]) or (keys[IR_COL["date"]] if len(keys) > IR_COL["date"] else None)
    for r in rows:
        day = date_key(r.get(c_date))
        if not day:
            continue
        add(assigning, user_id(r.get(c_assigned)), day, 1)
        add(tkt, user_id(r.get(c_added)), day, 1)
    return assigning, tkt


def parse_final_schedule(df):
    """شيت Final بعد الماكرو: ID | Date | Duration | Agent Name | TL | Code"""
    for i in range(min(len(df), 40)):
        cells = [norm(c) for c in df.iloc[i].tolist()]
        if not all(x in cells for x in ("id", "date", "duration")):
            continue
        c_id, c_date, c_dur = cells.index("id"), cells.index("date"), cells.index("duration")
        c_code = cells.index("code") if "code" in cells else -1
        index, days, code_totals = {}, set(), {}
        for row in df.iloc[i + 1:].itertuples(index=False):
            rid, day = text(row[c_id]), date_key(row[c_date])
            if not rid or not day:
                continue
            secs = parse_seconds(row[c_dur])
            e = index.setdefault(lookup_key(rid, day), {"total": 0, "has_total": False, "codes": {}})
            days.add(day)
            code = text(row[c_code]) if c_code >= 0 else ""
            if not code:
                e["total"] += secs
                e["has_total"] = True
            else:
                lower = code.lower()
                e["codes"][lower] = e["codes"].get(lower, 0) + secs
                code_totals[lower] = code_totals.get(lower, 0) + secs
        if index:
            return index, days, code_totals
    return None


def parse_schedule():
    """يرجّع (index, days, code_totals). index[key] = {total, has_total, codes{lower: secs}}"""
    path = find_data_file("Schedule")
    index, days, code_totals = {}, set(), {}
    if not path:
        return index, days, code_totals

    for _, df in read_sheets(path, header=None).items():
        final = parse_final_schedule(df)
        if final:
            return final

        # التقرير الخام: نحدد عمود Scheduled time وأعمدة العناوين من صف العناوين
        duration_col, label_end = SCHEDULE_COL["duration"], 4
        rows = df.values.tolist()
        for row in rows[:80]:
            cells = [norm(c) for c in row]
            hit = [i for i, c in enumerate(cells) if re.fullmatch(r"scheduledtime(hhmm)?", c)]
            if hit:
                duration_col = hit[0]
                first_val = [i for i, c in enumerate(cells) if re.match(r"(contracttime|worktime|paidtime)", c)]
                label_end = first_val[0] if first_val and first_val[0] > 1 else min(duration_col, 4)
                break

        login, day, agent_depth = "", "", -1
        for row in rows:
            row = list(row) + [None] * max(0, duration_col + 1 - len(row))
            depth, label = -1, None
            for i in range(label_end):
                if text(row[i]):
                    depth, label = i, row[i]
                    break
            if depth < 0:
                continue
            if norm(label) in {norm(w) for w in TOTALS_WORDS}:
                login, day, agent_depth = "", "", -1
                continue
            label_day = date_key(label)
            m = None if label_day else re.search(r"\b(\d{5,6})\b", text(label))
            if m:  # صف موظف
                login, agent_depth, day = m.group(1), depth, ""
                continue
            if not login or depth <= agent_depth:  # فريق/عنوان
                login, day, agent_depth = "", "", -1
                continue
            raw = row[duration_col]
            has_time = not is_blank(raw) and str(raw).strip() != ""
            if label_day:
                day = label_day
                days.add(label_day)
                e = index.setdefault(lookup_key(login, label_day), {"total": 0, "has_total": False, "codes": {}})
                if has_time:
                    e["total"] += parse_seconds(raw)
                    e["has_total"] = True
                continue
            if not day or not has_time:
                continue
            secs = parse_seconds(raw)
            if not secs:
                continue
            lower = text(label).lower()
            e = index.setdefault(lookup_key(login, day), {"total": 0, "has_total": False, "codes": {}})
            e["codes"][lower] = e["codes"].get(lower, 0) + secs
            code_totals[lower] = code_totals.get(lower, 0) + secs
        if index:
            break
    return index, days, code_totals


def schedule_seconds(entry, codes):
    if codes:
        return sum(entry["codes"].get(c, 0) for c in codes)
    return entry["total"] if entry["has_total"] else sum(entry["codes"].values())


# ─── الحساب ───────────────────────────────────────────────────────────────
def load_config():
    cfg = dict(DEFAULT_CONFIG)
    p = ROOT / "config.json"
    if p.exists():
        cfg.update(json.loads(p.read_text(encoding="utf-8")))
    return cfg


def build(structure, utl, ir, comp, sched_index, sched_days, codes, cfg):
    days = sorted({d for d in (date_key(val(r, "utlDate")) for r in utl) if d})
    if not days:
        days = set(sched_days)
        for r in ir:
            days.add(date_key(val(r, "irDate")))
        for r in comp:
            days.add(date_key(val(r, "compDate")))
        for r in structure:
            days.add(date_key(val(r, "structureDate")))
        days = sorted(d for d in days if d)

    ir_assigning, ir_tkt = ir_indexes(ir)
    talk = talk_index(utl)
    st_dur = sum_index(structure, "structureId", "structureDate", "structureDuration")
    comp_by_user = sum_index(comp, "compUser", "compDate", "compDuration")
    comp_by_id = sum_index(comp, "compId", "compDate", "compDuration")

    long_rows = []
    for r in structure:
        a = {k: text(val(r, k)) for k in
             ("structureId", "loginId", "perm", "ttsUser", "bssUser", "group", "agentName", "status", "tlId", "tlName")}
        a["status"] = a["status"] or "Active"
        if not a["agentName"] and not a["loginId"]:
            continue
        active = a["status"].lower() == "active"

        for day in days:
            assigning = lookup_first(ir_assigning, [user_id(a["ttsUser"])], day)
            tkt = lookup_first(ir_tkt, [user_id(a["ttsUser"])], day)
            system = round(tkt * cfg["secondsPerTicket"])
            talk_s = lookup_first(talk, [a["loginId"]], day)
            ck = lookup_key(user_id(a["ttsUser"]), day)
            comp_s = comp_by_user[ck] if ck and ck in comp_by_user else lookup_first(comp_by_id, [a["structureId"]], day)

            sch = None
            for cand in (a["structureId"], a["loginId"], a["agentName"]):
                e = sched_index.get(lookup_key(cand, day))
                if e:
                    sch = schedule_seconds(e, codes)
                    break
            if sch is None:
                sch = lookup_first(st_dur, [a["structureId"]], day)

            if sched_index and sch == 0:
                assigning = 0  # IF(Tele-SCH=0, 0, COUNTIFS(...))
            sch = round(sch * cfg["teleSchFactor"])  # Tele-SCH بعد الـ 90%
            loss = max(0, sch - (system + talk_s + comp_s))
            long_rows.append({
                "Teleopti ID": a["structureId"], "Login ID": a["loginId"], "Perm": a["perm"],
                "TTS User": a["ttsUser"], "BSS User": a["bssUser"], "Group": a["group"],
                "Agent Name": a["agentName"], "Status": a["status"], "TL ID": a["tlId"], "TL Name": a["tlName"],
                "Date": day, "Assigning Tkts": assigning, "TKT": tkt, "System": fmt_time(system),
                "Talk Time": fmt_time(talk_s), "Tele-SCH": fmt_time(sch), "Comp": fmt_time(comp_s),
                "Loss Time": fmt_time(loss) if active else a["status"],
            })
    return days, long_rows


EXPORT_FIXED = [("Teleopti ID", 12), ("Login ID", 11), ("Agent Name", 38), ("TTS User", 22), ("TL Name", 18), ("Status", 12)]
METRICS = ["Assigning Tkts", "TKT", "System", "Talk Time", "Tele-SCH", "Comp", "Loss Time"]


def write_matrix_sheet(ws, long_rows, days):
    """نفس شكل الإكسيل: صف 3 تواريخ مدموجة، صف 4 عناوين، 7 أعمدة لكل يوم (Outline) واليوم الأول مفتوح."""
    purple = PatternFill("solid", fgColor="7030A0")
    orange = PatternFill("solid", fgColor="E46C0A")
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    white_bold = Font(bold=True, color="FFFFFF")
    n_fixed, n_met = len(EXPORT_FIXED), len(METRICS)
    time_re = re.compile(r"^\d+:\d{2}:\d{2}$")

    ws.sheet_properties.outlinePr.summaryRight = True
    ws.freeze_panes = ws.cell(row=5, column=n_fixed + 1)

    for i, (label, width) in enumerate(EXPORT_FIXED, start=1):
        c = ws.cell(row=4, column=i, value=label)
        c.font, c.fill, c.alignment, c.border = white_bold, purple, center, border
        ws.column_dimensions[get_column_letter(i)].width = width

    for d, day in enumerate(days):
        start = n_fixed + d * n_met + 1
        dt_ = datetime.strptime(day, "%Y-%m-%d")
        ws.merge_cells(start_row=3, start_column=start, end_row=3, end_column=start + n_met - 1)
        h = ws.cell(row=3, column=start, value=f"{dt_.day}-{dt_.strftime('%b')}")
        h.font, h.alignment = Font(bold=True), center
        for c in range(start, start + n_met):
            ws.cell(row=3, column=c).border = border
        for m, label in enumerate(METRICS):
            c = ws.cell(row=4, column=start + m, value=label)
            c.font, c.fill, c.alignment, c.border = white_bold, orange if m == 0 else purple, center, border
            ws.column_dimensions[get_column_letter(start + m)].width = 14 if m == 0 else 11
        ws.column_dimensions.group(get_column_letter(start), get_column_letter(start + n_met - 2),
                                   outline_level=1, hidden=(d != 0))

    def num(txt):
        return parse_seconds(txt) / 86400

    per_agent = len(days)
    for r_i in range(0, len(long_rows), per_agent):
        chunk = long_rows[r_i:r_i + per_agent]
        row_no = 5 + r_i // per_agent
        first = chunk[0]
        for i, (label, _) in enumerate(EXPORT_FIXED, start=1):
            v = first[label]
            c = ws.cell(row=row_no, column=i, value=int(v) if re.fullmatch(r"\d+", str(v)) else v)
            c.alignment, c.border = center, border
        for d, item in enumerate(chunk):
            start = n_fixed + d * n_met + 1
            loss_is_time = bool(time_re.match(str(item["Loss Time"])))
            vals = [item["Assigning Tkts"], item["TKT"], num(item["System"]), num(item["Talk Time"]),
                    num(item["Tele-SCH"]), num(item["Comp"]),
                    num(item["Loss Time"]) if loss_is_time else item["Loss Time"]]
            for m, v in enumerate(vals):
                c = ws.cell(row=row_no, column=start + m, value=v)
                c.alignment, c.border = center, border
                if m >= 2 and (m < n_met - 1 or loss_is_time):
                    c.number_format = "[h]:mm:ss"
                if m == n_met - 1:
                    bad = loss_is_time and v > 0
                    c.fill = red if bad else green
                    c.font = Font(bold=True, color="9C0006" if bad else "006100")


def process_pipeline():
    cfg = load_config()
    structure = read_structure()
    if not structure:
        print("⚠️  لا يوجد ملف Structure (STR Loss.xlsx) صالح - تم الإيقاف")
        return 0

    utl = read_rows("UTL", ("utl", "log"), [ALIASES["utlUser"], ALIASES["utlDate"]])
    ir = read_rows("IR", ("ir", "ticket"), [ALIASES["irAssigned"] + ALIASES["irAdded"], ["added_on"]])
    comp = read_rows("Compensation", ("comp",), [ALIASES["compUser"] + ALIASES["compId"], ALIASES["compDuration"]])
    sched_index, sched_days, code_totals = parse_schedule()
    codes = {str(c).strip().lower() for c in cfg.get("scheduleCodes", []) if str(c).strip()}

    print(f"Structure={len(structure)} UTL={len(utl)} IR={len(ir)} Comp={len(comp)} "
          f"Schedule(agent-days)={len(sched_index)} Codes={sorted(codes) or 'إجمالي اليوم'}")
    if code_totals:
        print("الأكواد الموجودة في Schedule:", ", ".join(sorted(code_totals)))

    days, long_rows = build(structure, utl, ir, comp, sched_index, sched_days, codes, cfg)
    if not days or not long_rows:
        print("⚠️  لا توجد تواريخ داخل ملفات data/ - لا يوجد ما يُحدَّث")
        return 0

    long_df = pd.DataFrame(long_rows).fillna("")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    write_matrix_sheet(wb.active, long_df.to_dict("records"), days)
    wb.active.title = "Matrix"
    ws_long = wb.create_sheet("Daily_Long")
    ws_long.append(list(long_df.columns))
    for rec in long_df.itertuples(index=False):
        ws_long.append(list(rec))
    wb.save(OUTPUT)
    print(f"✅ تم حفظ التقرير في: {OUTPUT}  ({len(days)} يوم)")
    return 0


if __name__ == "__main__":
    sys.exit(process_pipeline())
